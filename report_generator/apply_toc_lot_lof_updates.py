import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt
import os

# Exact mappings derived from Microsoft Word rendering:
TOC_UPDATES = {
    # Front matter
    1: ("ACKNOWLEDGEMENT", "iii"),
    2: ("ABSTRACT", "iv"),
    3: ("LIST OF TABLES", "viii"),
    4: ("LIST OF FIGURES", "ix"),
    5: ("LIST OF ABBREVIATIONS", "x"),
    # Chapter 1
    6: ("INTRODUCTION", "1"),
    7: ("BACKGROUND OF THE WORK", "1"),
    8: ("Challenges in the present academic record system", "2"),
    9: ("MOTIVATION", "3"),
    10: ("Scope of the proposed work", "3"),
    # Chapter 2
    11: ("LITERATURE SURVEY", "5"),
    12: ("EDUCATIONAL DATA MINING AND STUDENT RISK MODELLING", "5"),
    13: ("MACHINE LEARNING CLASSIFIERS FOR ACADEMIC PERFORMANCE", "5"),
    14: ("NATURAL LANGUAGE INTERFACES AND LANGUAGE MODELS FOR ANALYTICS", "6"),
    15: ("INTERACTIVE DASHBOARDS FOR EDUCATIONAL DATA", "6"),
    16: ("GAP IDENTIFICATION", "7"),
    17: ("PROBLEM STATEMENT", "7"),
    # Chapter 3
    18: ("OBJECTIVES AND METHODOLOGY", "9"),
    19: ("OBJECTIVES OF THE PROPOSED WORK", "9"),
    20: ("FLOW DIAGRAM OF THE PROPOSED WORK", "10"),
    21: ("SELECTION OF DATA, TECHNIQUES, TOOLS AND TESTING METHODS", "13"),
    22: ("Academic data and normalization", "13"),
    23: ("Risk prediction technique", "15"),
    24: ("Conversational query technique", "16"),
    25: ("Visual analytics technique", "17"),
    26: ("Tools and technologies", "17"),
    27: ("Testing methods", "18"),
    # Chapter 4
    28: ("PROPOSED WORK MODULES", "19"),
    29: ("PROPOSED WORK", "19"),
    30: ("Academic record aggregator module", "19"),
    31: ("Risk prediction engine module", "20"),
    32: ("Conversational query and failover module", "20"),
    33: ("Analytics dashboard module", "21"),
    34: ("METHODOLOGY OF THE PROPOSED WORK", "22"),
    # Chapter 5
    35: ("RESULTS AND DISCUSSION", "24"),
    36: ("RESULTS", "24"),
    37: ("Data consolidation results", "24"),
    38: ("Risk prediction results", "25"),
    39: ("Conversational query results", "26"),
    40: ("Functional test results", "27"),
    41: ("Comparison with related published work", "28"),
    42: ("SIGNIFICANCE, STRENGTHS AND LIMITATIONS", "29"),
    43: ("Significance", "29"),
    44: ("Strengths", "29"),
    45: ("Limitations", "30"),
    46: ("COST BENEFIT ANALYSIS", "30"),
    # Chapter 6
    47: ("CONCLUSIONS AND FUTURE WORK", "32"),
    48: ("CONCLUSION", "32"),
    49: ("SUGGESTIONS FOR FUTURE WORK", "32"),
    # References & Appendices
    50: ("REFERENCES", "34"),
    51: ("APPENDICES", "36"),
    52: ("BILL OF MATERIALS", "36"),
    53: ("CODING", "37"),
    54: ("COMPETITION CERTIFICATE", "45"),
    55: ("COMPETITION PROOF", "46"),
    56: ("WORK CONTRIBUTION", "48"),
    57: ("PLAGIARISM REPORT", "49")
}

LOT_UPDATES = {
    1: ("3.1", "10"),
    2: ("3.2", "14"),
    3: ("3.3", "14"),
    4: ("3.4", "15"),
    5: ("3.5", "17"),
    6: ("4.1", "19"),
    7: ("4.2", "21"),
    8: ("4.3", "22"),
    9: ("5.1", "28"),
    10: ("5.2", "28"),
    11: ("5.3", "30"),
    12: ("A.1", "36")
}

LOF_UPDATES = {
    1: ("3.1", "11"),
    2: ("3.2", "13"),
    3: ("3.3", "16"),
    4: ("5.1", "25"),
    5: ("5.2", "26"),
    6: ("5.3", "27"),
    7: ("5.4", "27")
}

def set_cell_text(cell, new_text, align=WD_ALIGN_PARAGRAPH.CENTER, font_size_pt=12.0, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if not cell.paragraphs:
        p = cell.add_paragraph()
    else:
        p = cell.paragraphs[0]
        # remove extra paragraphs in cell if any
        for extra_p in cell.paragraphs[1:]:
            p._p.getparent().remove(extra_p._p)
            
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    if len(p.runs) > 0:
        p.runs[0].text = new_text
        p.runs[0].font.name = "Times New Roman"
        p.runs[0].font.size = Pt(font_size_pt)
        p.runs[0].font.bold = bold
        for extra_r in p.runs[1:]:
            extra_r.text = ""
    else:
        run = p.add_run(new_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size_pt)
        run.font.bold = bold

def update_report(doc_path):
    print(f"Loading {doc_path}...")
    doc = docx.Document(doc_path)
    
    # 1. Update Table 3: Table of Contents
    t_toc = doc.tables[3]
    print(f"\n--- Updating Table of Contents (Table 3, {len(t_toc.rows)} rows) ---")
    for row_idx, (expected_topic, new_page) in TOC_UPDATES.items():
        row = t_toc.rows[row_idx]
        title_cell = row.cells[1]
        page_cell = row.cells[2]
        
        # Check topic update for competition certificate & proof
        if row_idx == 54:
            set_cell_text(title_cell, "COMPETITION CERTIFICATE", align=WD_ALIGN_PARAGRAPH.LEFT, font_size_pt=12.0, bold=True)
            print("  Updated row 54 title: 'COMPETITION CERTIFICATE'")
        elif row_idx == 55:
            set_cell_text(title_cell, "COMPETITION PROOF", align=WD_ALIGN_PARAGRAPH.LEFT, font_size_pt=12.0, bold=True)
            print("  Updated row 55 title: 'COMPETITION PROOF'")

        old_page = page_cell.text.strip()
        set_cell_text(page_cell, new_page, align=WD_ALIGN_PARAGRAPH.CENTER, font_size_pt=12.0, bold=False)
        print(f"  TOC Row {row_idx:2d} | '{expected_topic[:40]}': {old_page} -> {new_page}")

    # 2. Update Table 4: List of Tables
    t_lot = doc.tables[4]
    print(f"\n--- Updating List of Tables (Table 4, {len(t_lot.rows)} rows) ---")
    for row_idx, (tbl_num, new_page) in LOT_UPDATES.items():
        row = t_lot.rows[row_idx]
        page_cell = row.cells[2]
        old_page = page_cell.text.strip()
        set_cell_text(page_cell, new_page, align=WD_ALIGN_PARAGRAPH.CENTER, font_size_pt=12.0, bold=False)
        print(f"  LOT Row {row_idx:2d} | Table {tbl_num}: {old_page} -> {new_page}")

    # 3. Update Table 5: List of Figures
    t_lof = doc.tables[5]
    print(f"\n--- Updating List of Figures (Table 5, {len(t_lof.rows)} rows) ---")
    for row_idx, (fig_num, new_page) in LOF_UPDATES.items():
        row = t_lof.rows[row_idx]
        page_cell = row.cells[2]
        old_page = page_cell.text.strip()
        set_cell_text(page_cell, new_page, align=WD_ALIGN_PARAGRAPH.CENTER, font_size_pt=12.0, bold=False)
        print(f"  LOF Row {row_idx:2d} | Figure {fig_num}: {old_page} -> {new_page}")

    doc.save(doc_path)
    print(f"\nSuccessfully updated all topics and page numbers in {doc_path}!")

if __name__ == "__main__":
    update_report("Logith Project Report (Corrected).docx")
