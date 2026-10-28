import docx

TOC_UPDATES = {
    1: ("ACKNOWLEDGEMENT", "iii"),
    2: ("ABSTRACT", "iv"),
    3: ("TABLE OF CONTENT", "vi"),
    4: ("LIST OF FIGURES", "viii"),
    5: ("LIST OF TABLES", "ix"),
    6: ("INTRODUCTION", "1"),
    7: ("Background of the Work", "1"),
    8: ("Problem Statement", "2"),
    9: ("Motivation and Scope of the Proposed Work", "3"),
    10: ("LITERATURE SURVEY", "6"),
    11: ("Educational Data Mining and Student Risk Modeling", "6"),
    12: ("Machine Learning Classifiers for Academic Progression", "6"),
    13: ("Natural Language Interfaces and LLMs for Institutional Analytics", "7"),
    14: ("Modern Web Dashboards and Interactive Visualizations in Education", "7"),
    15: ("Gap Identification and Problem Formulation", "8"),
    16: ("OBJECTIVES AND METHODOLOGY", "9"),
    17: ("Objectives of the Proposed Work", "9"),
    18: ("System Architecture Overview", "9"),
    19: ("Academic Data Ingestion & Normalization Pipeline", "10"),
    20: ("Machine Learning Risk Prediction Engine", "12"),
    21: ("Conversational AI Query Subsystem & Graph Specification Engine", "14"),
    22: ("Dynamic Frontend Visual Analytics Subsystem", "14"),
    23: ("Selection of Tools and Technologies", "15"),
    24: ("PROPOSED WORK MODULES", "17"),
    25: ("Ingestion & Academic Record Aggregator Module Implementation", "17"),
    26: ("Machine Learning Risk Prediction Engine Implementation", "18"),
    27: ("Conversational AI & Multi-Key Failover Engine Implementation", "18"),
    28: ("Dynamic Visual Analytics & Dashboard Implementation", "20"),
    29: ("RESULTS AND DISCUSSION", "23"),
    30: ("System Functionality Results", "23"),
    31: ("Significance, Strengths, and Limitations", "25"),
    32: ("Cost-Benefit Analysis", "26"),
    33: ("CONCLUSIONS AND FUTURE WORK", "28"),
    34: ("Conclusion", "28"),
    35: ("Suggestions for Future Work", "28"),
    36: ("REFERENCES", "30"),
    37: ("APPENDIX", "32"),
}

def update_toc(doc_path):
    print(f"Loading {doc_path}...")
    doc = docx.Document(doc_path)
    t = doc.tables[2]
    
    print(f"Table 2 rows count: {len(t.rows)}")
    
    changes = 0
    for row_idx, (expected_title, new_page) in TOC_UPDATES.items():
        row = t.rows[row_idx]
        actual_title = row.cells[1].text.strip().replace("\n", " ")
        if expected_title not in actual_title and actual_title not in expected_title:
            raise ValueError(f"Mismatch at row {row_idx}: expected '{expected_title}', found '{actual_title}'")
        
        page_cell = row.cells[2]
        p = page_cell.paragraphs[0]
        old_page = p.text.strip()
        
        if len(p.runs) > 0:
            p.runs[0].text = new_page
            # clear any extra runs if present
            for extra_r in p.runs[1:]:
                extra_r.text = ""
        else:
            p.text = new_page
            
        print(f"Row {row_idx:2d} | '{expected_title[:45]}': {old_page} -> {new_page}")
        changes += 1

    print(f"Successfully updated {changes} TOC page numbers in Table 2.")
    doc.save(doc_path)
    print(f"Saved {doc_path}.")

if __name__ == "__main__":
    update_toc("temp_local.docx")
