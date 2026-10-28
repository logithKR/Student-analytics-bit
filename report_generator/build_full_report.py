import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from front_matter_data import (
    PROJECT_TITLE, ABSTRACT_TEXT, TOC_ROWS,
    LIST_OF_FIGURES_ROWS, LIST_OF_TABLES_ROWS
)
from chapters_data_part1 import (
    CHAPTER_1_TITLE, CHAPTER_1_SECTIONS,
    CHAPTER_2_TITLE, CHAPTER_2_SECTIONS
)
from chapters_data_part2 import (
    CHAPTER_3_TITLE, CHAPTER_3_SECTIONS, TABLE_3_1,
    CHAPTER_3_SECTIONS_CONT, TABLE_3_2,
    CHAPTER_3_SECTIONS_CONT2, TABLE_3_3,
    CHAPTER_3_SECTIONS_CONT3, TABLE_3_4,
    CHAPTER_4_TITLE, CHAPTER_4_SECTIONS, TABLE_4_1,
    CHAPTER_4_SECTIONS_CONT, TABLE_4_2,
    CHAPTER_4_SECTIONS_CONT2, TABLE_4_3
)
from chapters_data_part3 import (
    CHAPTER_5_TITLE, CHAPTER_5_SECTIONS, TABLE_5_1,
    CHAPTER_5_SECTIONS_CONT, TABLE_5_2,
    CHAPTER_6_TITLE, CHAPTER_6_SECTIONS,
    REFERENCES, APPENDIX_I_BOM, WORK_CONTRIBUTION_TEXT
)

def build_report():
    print("Loading repaired template document...")
    template_path = "template_repaired.docx"
    if not os.path.exists(template_path):
        source_template = "Automated_Timetable_Generator_Report_Final (1).docx"
        if os.path.exists(source_template):
            import zipfile
            orig_update_crc = zipfile.ZipExtFile._update_crc
            zipfile.ZipExtFile._update_crc = lambda self, newdata: None
            with zipfile.ZipFile(source_template, 'r') as src:
                with zipfile.ZipFile(template_path, 'w', compression=zipfile.ZIP_DEFLATED) as dst:
                    for item in src.infolist():
                        data = src.read(item.filename)
                        dst.writestr(item.filename, data)
            zipfile.ZipExtFile._update_crc = orig_update_crc
            print("Successfully extracted and repaired template_repaired.docx")
    doc = docx.Document(template_path)

    # Helper styling functions
    def style_body_p(p, text, first_indent=True):
        p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Inches(0.5) if first_indent else Inches(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)

    def style_chapter_heading(p_num, num_text, p_title, title_text):
        p_num.text = num_text
        p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_num.paragraph_format.line_spacing = 1.5
        p_num.paragraph_format.first_line_indent = Inches(0)
        p_num.paragraph_format.space_before = Pt(24)
        p_num.paragraph_format.space_after = Pt(6)
        for r in p_num.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(14)
            r.font.bold = True

        p_title.text = title_text
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.line_spacing = 1.5
        p_title.paragraph_format.first_line_indent = Inches(0)
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(18)
        for r in p_title.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(14)
            r.font.bold = True

    def style_section_heading(p, text):
        p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            r.font.bold = True

    def style_caption(p, text):
        p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(12)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.bold = True

    def add_custom_table(table_spec):
        headers = table_spec["headers"]
        rows = table_spec["rows"]
        caption = table_spec["caption"]

        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Apply borders
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

        # Populate headers
        for col_idx, header_text in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = header_text
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="D9D9D9"/>')
            tcPr.append(shd)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)
                r.font.bold = True

        # Populate rows
        for row_idx, row_data in enumerate(rows):
            for col_idx, cell_value in enumerate(row_data):
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = str(cell_value)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10.5)

        # Add caption below table
        p_cap = doc.add_paragraph()
        style_caption(p_cap, caption)
        return table

    def add_custom_figure(image_path, caption_text, width_inches=5.8):
        if not os.path.exists(image_path):
            print(f"[!] Warning: Image {image_path} not found.")
            return
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(12)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        style_caption(p_cap, caption_text)

    # 1. UPDATE FRONT MATTER
    print("Updating Front Matter text...")
    # Cover Page Title
    p0 = doc.paragraphs[0]
    p0.text = PROJECT_TITLE
    for r in p0.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(18)
        r.font.bold = True

    # Bonafide Certificate Title
    p27 = doc.paragraphs[27]
    p27.text = f'Certified that this project report “{PROJECT_TITLE}” is the Bonafide work of “LOGITHKUMAR K R (7376231CD124), RAMESH S (7376231CD141)” who carried out the project work under my supervision.'
    for r in p27.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True

    # Declaration Title
    for p_idx in [52, 53, 54]:
        p_decl = doc.paragraphs[p_idx]
        if "We affirm that" in p_decl.text:
            p_decl.text = f'We affirm that the project work titled “{PROJECT_TITLE}” being submitted in partial fulfilment for the award of the degree of Bachelor of Engineering in Computer Science and Design, is the record of original work done by us under the guidance of Ms. INDHUMATHI K, Assistant Professor, Department of Computer Science and Design. It has not formed a part of any other project work(s) submitted for the award of any degree or diploma, either in this or any other University.'
            for r in p_decl.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.font.bold = True
            break

    # Abstract Text (Find abstract paragraphs)
    abs_indices = []
    for i in range(85, 98):
        p = doc.paragraphs[i]
        if "Academic timetabling" in p.text or "This project presents" in p.text or "The backend engine" in p.text or "Additionally, a Conflict" in p.text:
            abs_indices.append(i)

    print(f"Replacing {len(abs_indices)} abstract paragraphs with new project abstract...")
    for idx, new_abs_text in zip(abs_indices, ABSTRACT_TEXT):
        p = doc.paragraphs[idx]
        p.text = new_abs_text
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Inches(0.5)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)

    # 2. UPDATE TABLES (TOC, List of Figures, List of Tables)
    print("Updating Table of Contents (Table 2)...")
    t_toc = doc.tables[2]
    for row_idx, data in enumerate(TOC_ROWS):
        if row_idx < len(t_toc.rows):
            row = t_toc.rows[row_idx]
            for col_idx, val in enumerate(data):
                cell = row.cells[col_idx]
                cell.text = val
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(11)
                    if row_idx == 0 or (col_idx == 1 and val.isupper() and len(val) > 2):
                        r.font.bold = True

    print("Updating List of Figures (Table 3)...")
    t_lof = doc.tables[3]
    for row_idx, data in enumerate(LIST_OF_FIGURES_ROWS):
        if row_idx < len(t_lof.rows):
            row = t_lof.rows[row_idx]
            for col_idx, val in enumerate(data):
                cell = row.cells[col_idx]
                cell.text = val
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(11)
                    if row_idx == 0:
                        r.font.bold = True

    print("Updating List of Tables (Table 4)...")
    t_lot = doc.tables[4]
    for row_idx, data in enumerate(LIST_OF_TABLES_ROWS):
        if row_idx < len(t_lot.rows):
            row = t_lot.rows[row_idx]
            for col_idx, val in enumerate(data):
                cell = row.cells[col_idx]
                cell.text = val
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(11)
                    if row_idx == 0:
                        r.font.bold = True

    # 3. SPLICE BODY ELEMENTS & PRESERVE PLAGIARISM SECTION
    print("Splicing body elements...")
    body = doc._body._element

    # Find plagiarism start index
    plag_start_idx = None
    for i, c in enumerate(body):
        txt = ''.join(c.itertext()).strip()
        if "III. PLAGIARISM REPORT" in txt:
            plag_start_idx = i
            break

    if plag_start_idx is None:
        plag_start_idx = 422
    print(f"Plagiarism section starts at body index {plag_start_idx}")

    plagiarism_elements = [body[i] for i in range(plag_start_idx, len(body) - 1)]
    print(f"Extracted {len(plagiarism_elements)} plagiarism elements to preserve.")

    # Remove elements from 109 to plag_start_idx + len(plagiarism_elements)
    remove_count = (len(body) - 1) - 109
    for _ in range(remove_count):
        body.remove(body[109])

    print(f"Cleared old body content. Remaining body elements: {len(body)}")

    # 4. APPEND NEW BODY CHAPTERS & SECTIONS
    print("Generating Chapter I...")
    p_num = doc.add_paragraph()
    p_title = doc.add_paragraph()
    style_chapter_heading(p_num, "CHAPTER I", p_title, CHAPTER_1_TITLE)
    for sec in CHAPTER_1_SECTIONS:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    doc.add_page_break()

    print("Generating Chapter II...")
    p_num = doc.add_paragraph()
    p_title = doc.add_paragraph()
    style_chapter_heading(p_num, "CHAPTER II", p_title, CHAPTER_2_TITLE)
    for sec in CHAPTER_2_SECTIONS:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    doc.add_page_break()

    print("Generating Chapter III...")
    p_num = doc.add_paragraph()
    p_title = doc.add_paragraph()
    style_chapter_heading(p_num, "CHAPTER III", p_title, CHAPTER_3_TITLE)
    for sec in CHAPTER_3_SECTIONS:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    # Figure 3.1
    add_custom_figure("images/fig3_1_system_architecture.png", "Figure 3.1: Academic Intelligence Platform System Architecture Overview", 6.0)

    # Table 3.1
    add_custom_table(TABLE_3_1)

    for sec in CHAPTER_3_SECTIONS_CONT:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    # Table 3.2
    add_custom_table(TABLE_3_2)

    for sec in CHAPTER_3_SECTIONS_CONT2:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    # Table 3.3
    add_custom_table(TABLE_3_3)

    for sec in CHAPTER_3_SECTIONS_CONT3:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    # Figure 3.2
    add_custom_figure("images/fig3_2_ai_failover_flowchart.png", "Figure 3.2: End-to-End Conversational AI & Failover Execution Flow", 6.0)

    # Table 3.4
    add_custom_table(TABLE_3_4)

    doc.add_page_break()

    print("Generating Chapter IV...")
    p_num = doc.add_paragraph()
    p_title = doc.add_paragraph()
    style_chapter_heading(p_num, "CHAPTER IV", p_title, CHAPTER_4_TITLE)
    for sec in CHAPTER_4_SECTIONS:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    # Table 4.1
    add_custom_table(TABLE_4_1)

    for sec in CHAPTER_4_SECTIONS_CONT:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    # Table 4.2
    add_custom_table(TABLE_4_2)

    for sec in CHAPTER_4_SECTIONS_CONT2:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    # Table 4.3
    add_custom_table(TABLE_4_3)

    # Figures 4.1, 4.2, 4.3
    add_custom_figure("images/department_overview.jpeg", "Figure 4.1: Department Performance Overview Dashboard", 6.0)
    add_custom_figure("images/student_analytics.jpeg", "Figure 4.2: Individual Student Academic Analytics & Progression View", 6.0)
    add_custom_figure("images/ai_chat_assistant.jpeg", "Figure 4.3: Conversational AI Analytics & Dynamic Query Interface", 6.0)

    doc.add_page_break()

    print("Generating Chapter V...")
    p_num = doc.add_paragraph()
    p_title = doc.add_paragraph()
    style_chapter_heading(p_num, "CHAPTER V", p_title, CHAPTER_5_TITLE)
    for sec in CHAPTER_5_SECTIONS:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    # Table 5.1
    add_custom_table(TABLE_5_1)

    # Figures 5.1, 5.2
    add_custom_figure("images/hod_portal_mockup.jpeg", "Figure 5.1: Management HOD Analytics Portal Interface", 5.8)
    add_custom_figure("images/architecture_mockup.jpeg", "Figure 5.2: Academic Performance Analytics Hardware & Architecture Mockup", 5.8)

    for sec in CHAPTER_5_SECTIONS_CONT:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    # Table 5.2
    add_custom_table(TABLE_5_2)

    doc.add_page_break()

    print("Generating Chapter VI...")
    p_num = doc.add_paragraph()
    p_title = doc.add_paragraph()
    style_chapter_heading(p_num, "CHAPTER VI", p_title, CHAPTER_6_TITLE)
    for sec in CHAPTER_6_SECTIONS:
        p_sec = doc.add_paragraph()
        style_section_heading(p_sec, sec["heading"])
        for text in sec["paragraphs"]:
            p = doc.add_paragraph()
            style_body_p(p, text)

    doc.add_page_break()

    print("Generating References...")
    p_ref = doc.add_paragraph()
    p_ref.text = "REFERENCES"
    p_ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ref.paragraph_format.line_spacing = 1.5
    p_ref.paragraph_format.space_before = Pt(24)
    p_ref.paragraph_format.space_after = Pt(18)
    for r in p_ref.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.font.bold = True

    for ref_entry in REFERENCES:
        p = doc.add_paragraph()
        p.text = ref_entry
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Inches(-0.35)
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)

    doc.add_page_break()

    print("Generating Appendix I & II...")
    p_app = doc.add_paragraph()
    p_app.text = "APPENDIX"
    p_app.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_app.paragraph_format.line_spacing = 1.5
    p_app.paragraph_format.space_before = Pt(24)
    p_app.paragraph_format.space_after = Pt(18)
    for r in p_app.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.font.bold = True

    # Appendix I: Bill of Materials
    p_bom_head = doc.add_paragraph()
    p_bom_head.text = "I. BILL OF MATERIALS"
    p_bom_head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_bom_head.paragraph_format.space_before = Pt(14)
    p_bom_head.paragraph_format.space_after = Pt(6)
    for r in p_bom_head.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.bold = True

    p_bom_desc = doc.add_paragraph()
    style_body_p(p_bom_desc, "The Academic Intelligence Platform is a software-only analytics and machine learning solution requiring no specialized proprietary hardware components. All development and deployment were executed utilizing standard institutional computing infrastructure and open-source software frameworks. Table 14 outlines the complete software component inventory, open-source licensing models, and computational resource allocations.", True)

    add_custom_table(APPENDIX_I_BOM)

    # Appendix II: Work Contribution
    p_wc_head = doc.add_paragraph()
    p_wc_head.text = "II. WORK CONTRIBUTION"
    p_wc_head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_wc_head.paragraph_format.space_before = Pt(14)
    p_wc_head.paragraph_format.space_after = Pt(6)
    for r in p_wc_head.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.bold = True

    p_wc_desc = doc.add_paragraph()
    style_body_p(p_wc_desc, "The Academic Intelligence Platform was conceptualized, engineered, and evaluated through collaborative teamwork. The distinct technical contributions and architectural leadership of each student team member are documented below.", True)

    for member_title, member_desc in WORK_CONTRIBUTION_TEXT:
        p_m_title = doc.add_paragraph()
        p_m_title.text = member_title
        p_m_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_m_title.paragraph_format.space_before = Pt(8)
        p_m_title.paragraph_format.space_after = Pt(2)
        p_m_title.paragraph_format.first_line_indent = Inches(0)
        for r in p_m_title.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(12.5)
            r.font.bold = True

        p_m_desc = doc.add_paragraph()
        style_body_p(p_m_desc, member_desc, True)

    doc.add_page_break()

    # 5. RE-INSERT PRESERVED PLAGIARISM SECTION
    print("Re-inserting preserved Plagiarism report section...")
    sectPr = body[-1]
    for elem in plagiarism_elements:
        sectPr.addprevious(elem)

    # 6. SAVE FINAL DOCUMENT
    output_filename = "Project_Report_Final.docx"
    print(f"Saving finalized report to {output_filename}...")
    doc.save(output_filename)
    print("SUCCESS! Project_Report_Final.docx created successfully.")

if __name__ == "__main__":
    build_report()
